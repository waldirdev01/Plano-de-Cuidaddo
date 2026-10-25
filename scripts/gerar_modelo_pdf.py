"""Build physician materials from the same versioned text used by parser tests.

Run with the bundled workspace Python (python-docx and reportlab).
PDFs exported from DOCX must additionally pass render_docx.py and PDFKit checks.
"""
from pathlib import Path
import re
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'modelo-pdf-v1'
TMP = ROOT / 'tmp' / 'modelo-pdf-v1'
OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)


def criar_docx(source, destino):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    for node in doc.styles.element.iter(qn('w:pBdr')):
        node.getparent().remove(node)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.04
    doc.styles['Title'].font.size = Pt(24)
    doc.styles['Title'].paragraph_format.space_after = Pt(12)
    doc.styles['Heading 1'].font.size = Pt(14)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
    for raw in source.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line == 'Plano de Cuidado':
            doc.add_paragraph(line, 'Title')
        elif line in ['DADOS DO PLANO', 'ORIENTACOES LIVRES'] or re.fullmatch(r'(?:META|REFEICAO) [1-8]', line):
            p = doc.add_paragraph(line, 'Heading 1')
            if line in ['META 2', 'REFEICAO 3']:
                p.paragraph_format.page_break_before = True
        else:
            p = doc.add_paragraph()
            if ':' in line and not line.startswith('http'):
                label, value = line.split(':', 1)
                p.add_run(label + ':').bold = True
                run = p.add_run(value)
                if '[PREENCHER' in value:
                    run.italic = True
            else:
                run = p.add_run(line)
                if line.startswith('FIM '):
                    run.font.size = Pt(9)
    doc.core_properties.title = 'Plano de Cuidado padrão 1'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.save(destino)


def criar_guia():
    pdfmetrics.registerFont(TTFont('ArialLocal', '/System/Library/Fonts/Supplemental/Arial.ttf'))
    pdfmetrics.registerFont(TTFont('ArialLocalBold', '/System/Library/Fonts/Supplemental/Arial Bold.ttf'))
    body = ParagraphStyle('Body', fontName='ArialLocal', fontSize=10.5, leading=14, spaceAfter=8, textColor=colors.black, alignment=TA_LEFT)
    title = ParagraphStyle('Title', parent=body, fontName='ArialLocalBold', fontSize=21, leading=25, spaceAfter=14)
    heading = ParagraphStyle('Heading', parent=body, fontName='ArialLocalBold', fontSize=12, leading=15, spaceBefore=10, spaceAfter=7, keepWithNext=True)
    source = (ROOT / 'docs/modelos/orientacao-medica-padrao-1.md').read_text()
    story = []
    for chunk in source.strip().split('\n\n'):
        if chunk.startswith('# '):
            story.append(Paragraph(escape(chunk[2:]), title))
        elif chunk.startswith('## '):
            story.append(Paragraph(escape(chunk[3:]), heading))
        elif chunk.startswith('Quantidade diaria:'):
            story.append(Paragraph(escape(chunk).replace('\n', '<br/>'), body))
        else:
            story.append(Paragraph(escape(chunk).replace('\n', '<br/>'), body))
    SimpleDocTemplate(str(OUT / 'orientacao-para-medica.pdf'), pagesize=letter,
        rightMargin=52, leftMargin=52, topMargin=45, bottomMargin=45,
        title='Como preencher o plano para o aplicativo', author='').build(story)


if __name__ == '__main__':
    criar_docx((ROOT / 'docs/modelos/plano-cuidado-padrao-1.txt').read_text(), OUT / 'modelo-plano-cuidado-padrao-1.docx')
    criar_docx((ROOT / 'test/fixtures/plano-cuidado-padrao-1.txt').read_text(), TMP / 'exemplo-padrao-1.docx')
    criar_guia()
    print(OUT)
